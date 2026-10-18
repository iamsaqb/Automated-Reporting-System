
from pathlib import Path
import logging
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

# --------------------------------------------------
# 1. Project folders and file paths
# --------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input_files"
OUTPUT_DIR = BASE_DIR / "output_reports"
LOG_DIR = BASE_DIR / "logs"

INPUT_FILE = INPUT_DIR / "retail_sales_dataset.csv"
OUTPUT_FILE = OUTPUT_DIR / "Automated_Sales_Report.xlsx"
LOG_FILE = LOG_DIR / "report.log"

# Create required folders automatically
INPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
LOG_DIR.mkdir(parents=True, exist_ok=True)

# --------------------------------------------------
# 2. Logging setup
# --------------------------------------------------
logging.basicConfig(
    filename=str(LOG_FILE),
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    encoding="utf-8"
)


# --------------------------------------------------
# 3. Generate automated sales report
# --------------------------------------------------
def generate_report():
    print("Starting automated sales reporting...")
    logging.info("Automated sales reporting started.")

    try:
        # Load and validate data
        if not INPUT_FILE.exists():
            raise FileNotFoundError(
                f"CSV file not found: {INPUT_FILE}\n"
                "Place retail_sales_dataset.csv inside input_files."
            )

        df = pd.read_csv(INPUT_FILE)

        required_columns = [
            "Transaction ID",
            "Date",
            "Customer ID",
            "Gender",
            "Age",
            "Product Category",
            "Quantity",
            "Price per Unit",
            "Total Amount"
        ]

        missing = [
            col for col in required_columns
            if col not in df.columns
        ]

        if missing:
            raise ValueError(
                f"Missing required columns: {missing}"
            )

        if df.empty:
            raise ValueError("The input CSV file contains no data.")

        # Convert dates and remove invalid dates
        df["Date"] = pd.to_datetime(
            df["Date"],
            errors="coerce"
        )

        invalid_dates = df["Date"].isna().sum()

        if invalid_dates:
            logging.warning(
                "Removed %s rows with invalid dates.",
                invalid_dates
            )
            df = df.dropna(subset=["Date"])

        if df.empty:
            raise ValueError(
                "No valid records remain after date validation."
            )

        logging.info(
            "Dataset loaded successfully: %s rows.",
            len(df)
        )

        # --------------------------------------------------
        # 4. Sales summary
        # --------------------------------------------------
        total_revenue = df["Total Amount"].sum()
        total_transactions = df["Transaction ID"].nunique()
        total_units_sold = df["Quantity"].sum()
        average_transaction_value = df["Total Amount"].mean()

        sales_summary = pd.DataFrame({
            "Metric": [
                "Total Revenue",
                "Total Transactions",
                "Total Units Sold",
                "Average Transaction Value"
            ],
            "Value": [
                total_revenue,
                total_transactions,
                total_units_sold,
                round(average_transaction_value, 2)
            ]
        })

        # --------------------------------------------------
        # 5. Category analysis
        # --------------------------------------------------
        category_summary = (
            df.groupby("Product Category")["Total Amount"]
            .sum()
            .sort_values(ascending=False)
            .reset_index()
        )

        category_summary.columns = [
            "Product Category",
            "Total Revenue"
        ]

        # --------------------------------------------------
        # 6. Monthly analysis
        # --------------------------------------------------
        df["Month"] = df["Date"].dt.to_period("M").astype(str)

        monthly_summary = (
            df.groupby("Month")["Total Amount"]
            .sum()
            .reset_index()
            .sort_values("Month")
        )

        monthly_summary.columns = [
            "Month",
            "Total Revenue"
        ]

        # --------------------------------------------------
        # 7. Gender analysis
        # --------------------------------------------------
        gender_summary = (
            df.groupby("Gender")["Total Amount"]
            .agg(["sum", "count", "mean"])
            .reset_index()
        )

        gender_summary.columns = [
            "Gender",
            "Total Revenue",
            "Transactions",
            "Average Transaction Value"
        ]

        # --------------------------------------------------
        # 8. Write data to Excel
        # --------------------------------------------------
        with pd.ExcelWriter(
            OUTPUT_FILE,
            engine="openpyxl"
        ) as writer:
            sales_summary.to_excel(
                writer,
                sheet_name="Sales Summary",
                index=False
            )

            category_summary.to_excel(
                writer,
                sheet_name="Category Analysis",
                index=False
            )

            monthly_summary.to_excel(
                writer,
                sheet_name="Monthly Analysis",
                index=False
            )

            gender_summary.to_excel(
                writer,
                sheet_name="Gender Analysis",
                index=False
            )

        # --------------------------------------------------
        # 9. Format Excel workbook
        # --------------------------------------------------
        wb = load_workbook(OUTPUT_FILE)

        for ws in wb.worksheets:
            # Format header row
            for cell in ws[1]:
                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )
                cell.fill = PatternFill(
                    "solid",
                    fgColor="1F4E78"
                )
                cell.alignment = Alignment(
                    horizontal="center"
                )

            # Automatically adjust column widths
            for column_cells in ws.columns:
                max_length = max(
                    (
                        len(str(cell.value))
                        for cell in column_cells
                        if cell.value is not None
                    ),
                    default=0
                )

                ws.column_dimensions[
                    column_cells[0].column_letter
                ].width = min(max_length + 4, 35)

            # Keep headings visible when scrolling
            ws.freeze_panes = "A2"

            # Add filters to the data table
            ws.auto_filter.ref = ws.dimensions

        # --------------------------------------------------
        # 10. Category revenue bar chart
        # --------------------------------------------------
        ws = wb["Category Analysis"]

        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Revenue by Product Category"
        chart.x_axis.title = "Total Revenue"
        chart.y_axis.title = "Product Category"

        chart.add_data(
            Reference(
                ws,
                min_col=2,
                min_row=1,
                max_row=ws.max_row
            ),
            titles_from_data=True
        )

        chart.set_categories(
            Reference(
                ws,
                min_col=1,
                min_row=2,
                max_row=ws.max_row
            )
        )

        chart.height = 7
        chart.width = 14
        ws.add_chart(chart, "D2")

        # --------------------------------------------------
        # 11. Monthly revenue line chart
        # --------------------------------------------------
        ws = wb["Monthly Analysis"]

        chart = LineChart()
        chart.style = 13
        chart.title = "Monthly Revenue Trend"
        chart.x_axis.title = "Month"
        chart.y_axis.title = "Total Revenue"

        chart.add_data(
            Reference(
                ws,
                min_col=2,
                min_row=1,
                max_row=ws.max_row
            ),
            titles_from_data=True
        )

        chart.set_categories(
            Reference(
                ws,
                min_col=1,
                min_row=2,
                max_row=ws.max_row
            )
        )

        chart.height = 7
        chart.width = 14
        ws.add_chart(chart, "D2")

        # Save the completed workbook
        wb.save(OUTPUT_FILE)
        wb.close()

        print("Report generated successfully!")
        print(f"Saved to: {OUTPUT_FILE}")
        print(f"Log file: {LOG_FILE}")

        logging.info(
            "Report generated successfully: %s",
            OUTPUT_FILE
        )

    except Exception:
        logging.exception("Sales report generation failed.")
        print("Report generation failed. Check logs/report.log.")
        raise


# --------------------------------------------------
# 12. Run the program
# --------------------------------------------------
if __name__ == "__main__":
    generate_report()
